"""用真实 XLSX 序列化验证外部资产文本不会成为公式。"""
import ast
import copy
import importlib.util
from io import BytesIO
from pathlib import Path
import re
import sys
import types
import unittest
from unittest.mock import patch
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
# 使用应用已有的旧 openpyxl / 新 numpy 兼容处理，不导入数据库或路由。
import app
import openpyxl
from openpyxl.styles import Font

spec = importlib.util.spec_from_file_location('literal_excel', ROOT / 'app/utils/excel_export.py')
literal = importlib.util.module_from_spec(spec)
spec.loader.exec_module(literal)


def extract(path, names, env):
    nodes = ast.parse((ROOT / path).read_text()).body
    selected = [copy.deepcopy(n) for n in nodes if isinstance(n, (ast.FunctionDef, ast.ClassDef)) and n.name in names]
    exec(compile(ast.fix_missing_locations(ast.Module(body=selected, type_ignores=[])), str(path), 'exec'), env)
    return env


class ExcelTests(unittest.TestCase):
    def roundtrip(self, workbook):
        output = BytesIO()
        workbook.save(output)
        with ZipFile(output) as archive:
            for name in archive.namelist():
                if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                    self.assertNotIn(b'<f>', archive.read(name))
        return openpyxl.load_workbook(BytesIO(output.getvalue()), data_only=False)

    def test_literal_rows_preserve_text_numbers_and_empty_values(self):
        workbook = openpyxl.Workbook()
        values = ['=1+1', '+1+1', '-1+1', '@SUM(A1)', '普通标题', 200, -3, None]
        literal.append_literal_row(workbook.active, values)
        sheet = self.roundtrip(workbook).active
        self.assertEqual([sheet.cell(1, c).value for c in range(1, 9)], values)
        self.assertEqual(sheet['A1'].data_type, 's')
        self.assertEqual(sheet['F1'].data_type, 'n')

    def test_write_only_sheet_also_stores_formula_text_literally(self):
        workbook = openpyxl.Workbook(write_only=True)
        sheet = workbook.create_sheet()
        literal.append_literal_row(sheet, ['=1+1', 42])
        saved = self.roundtrip(workbook).active
        self.assertEqual(saved['A1'].value, '=1+1')
        self.assertEqual(saved['A1'].data_type, 's')
        self.assertEqual(saved['B1'].value, 42)

    def test_statistics_cell_writes_preserve_external_product_text(self):
        workbook = openpyxl.Workbook()
        literal.set_literal_cell(workbook.active, 5, 11, '=1+1')
        literal.set_literal_cell(workbook.active, 5, 12, 7)
        saved = self.roundtrip(workbook).active
        self.assertEqual(saved['K5'].value, '=1+1')
        self.assertEqual(saved['K5'].data_type, 's')
        self.assertEqual(saved['L5'].value, 7)

    def task_export(self, **extra):
        env = {'Workbook': openpyxl.Workbook, 'Font': Font, 're': re,
               'append_literal_row': literal.append_literal_row, 'set_literal_cell': literal.set_literal_cell}
        env.update(extra)
        return extract('app/routes/export.py', ['SaveTask'], env)['SaveTask']('fixture')

    def test_real_site_export_keeps_title_literal_and_status_numeric(self):
        exporter = self.task_export(get_site_data=lambda _: [{
            'site': 'https://alpha.example', 'title': '=1+1', 'finger': [], 'status': 200, 'favicon': {}}])
        exporter.build_site_xl()
        saved = self.roundtrip(exporter.wb).active
        self.assertEqual(saved['B2'].value, '=1+1')
        self.assertEqual(saved['B2'].data_type, 's')
        self.assertEqual(saved['D2'].value, 200)

    def test_real_service_export_keeps_product_literal(self):
        exporter = self.task_export(get_ip_data=lambda _: [{
            'ip': '192.0.2.1', 'port_info': [{'port_id': 80, 'service_name': 'http', 'product': '=1+1'}]}])
        exporter.build_service_xl()
        saved = self.roundtrip(exporter.wb)['系统服务']
        self.assertEqual(saved['D2'].value, '=1+1')
        self.assertEqual(saved['D2'].data_type, 's')

    def test_real_company_export_keeps_names_literal(self):
        env = {'append_literal_row': literal.append_literal_row, 're': re}
        extract('app/routes/icp.py', ['normalize_icp_asset_item', '_build_excel_response'], env)
        response_module = types.SimpleNamespace(make_response=lambda data: types.SimpleNamespace(data=data, headers={}))
        with patch.dict(sys.modules, {'flask': response_module}):
            response = env['_build_excel_response']([{'query_type': 'web', 'unitName': '=1+1',
                                                     'domain': 'alpha.example', 'serviceName': '=2+2'}])
        workbook = openpyxl.load_workbook(BytesIO(response.data), data_only=False)
        self.assertEqual(workbook['web']['A2'].value, '=1+1')
        self.assertEqual(workbook['web']['A2'].data_type, 's')
        self.assertEqual(workbook['web']['E2'].value, '=2+2')
        self.assertEqual(workbook['web']['E2'].data_type, 's')


if __name__ == '__main__':
    unittest.main()
