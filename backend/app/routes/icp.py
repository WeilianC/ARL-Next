from flask_restx import Namespace, fields
from app.utils import get_logger, auth, build_ret, curr_date, conn_db
from app.modules import TaskStatus, ErrorMsg
from app.utils.task_statistics import get_task_statistics, attach_task_statistics
from app.utils.excel_export import append_literal_row
from . import ARLResource, get_arl_parser, base_query_fields
import bson

ns = Namespace('icp', description="ICP 备案查询")
logger = get_logger()


def _request_recon(action, options):
    """只有微服务确认接收或停止后，调用方才报告成功。"""
    import requests
    response = requests.post(f"http://osint-service:16181/api/v1/recon/{action}", json=options, timeout=5)
    response.raise_for_status()
    result = response.json()
    if not isinstance(result, dict) or result.get("code") != 200 or result.get("task_id") != options.get("task_id"):
        message = result.get("msg", "微服务未确认操作") if isinstance(result, dict) else "微服务响应格式错误"
        raise RuntimeError(message)
    return result

# 基础搜索字段
base_search_icp_task_fields = {
    'name': fields.String(required=False, description="任务名"),
    'target': fields.String(description="查询目标"),
    'status': fields.String(description="任务状态"),
    'include_statistics': fields.Integer(description="设为 1 时附带全局企业查询状态统计"),
    'query_type': fields.String(description="查询类型"),
    'end_time__gte': fields.String(description="结束时间大于等于"),
    'end_time__lte': fields.String(description="结束时间小于等于"),
    'end_time__dgt': fields.String(description="结束时间大于等于(兼容)"),
    'end_time__dlt': fields.String(description="结束时间小于等于(兼容)"),
    'start_time__gte': fields.String(description="开始时间大于等于"),
    'start_time__lte': fields.String(description="开始时间小于等于"),
    '_id': fields.String(description="任务ID")
}
base_search_icp_task_fields.update(base_query_fields)
search_icp_task_fields = ns.model('SearchIcpTask', base_search_icp_task_fields)

# 新建任务参数
add_icp_task_fields = ns.model('AddIcpTask', {
    'name': fields.String(required=False, example="腾讯备案查询", description="任务名"),
    'target': fields.String(required=True, example="深圳市腾讯计算机系统有限公司", description="查询目标"),
    'query_type': fields.List(fields.String, required=False, example=["web", "app"], description="查询类型列表(web/app/mapp/kapp)"),
    'depth': fields.Integer(required=False, default=1, description="递归查询层数(1~20)"),
    'invest_ratio': fields.Float(required=False, default=50, description="对外投资最小比例(%)"),
    'enable_icp': fields.Boolean(required=False, default=True, description="是否联动工信部ICP双轨查重与补全"),
    'auto_group': fields.Boolean(required=False, default=False, description="测绘完成后是否由后端自动纳管建组"),
    'auto_group_scan': fields.Boolean(required=False, default=False, description="自动建组后是否立即下发主动探测"),
    'auto_group_task_type': fields.String(required=False, default="oneshot", description="自动探测任务类型(oneshot/periodic)"),
    'auto_group_policy_id': fields.String(required=False, default="", description="自动探测使用的扫描策略ID"),
    'auto_group_interval_hours': fields.Integer(required=False, default=24, description="周期探测运行间隔(小时)"),
    'group_id': fields.String(required=False, default="", description="建组时关联的所属集团ID")
})

add_tyc_task_fields = ns.model('AddTycTask', {
    'name': fields.String(required=False, description="任务名称(选填，留空默认企业全称或TYC_GID)"),
    'company_name': fields.String(required=False, description="企业官方工商全称(选填，留空自动根据GID定向解析)"),
    'gid': fields.String(required=True, description="公司/节点GID"),
    'depth': fields.Integer(required=False, default=1, description="递归查询层数"),
    'invest_ratio': fields.Float(required=False, default=50, description="对外投资最小比例(%)"),
    'query_type': fields.List(fields.String, required=True, description="查询类型(web/app/mapp/wechat/weibo/invest)"),
    'enable_icp': fields.Boolean(required=False, default=True, description="是否联动工信部ICP双轨查重与补全"),
    'auto_group': fields.Boolean(required=False, default=False, description="测绘完成后是否由后端自动纳管建组"),
    'auto_group_scan': fields.Boolean(required=False, default=False, description="自动建组后是否立即下发主动探测"),
    'auto_group_task_type': fields.String(required=False, default="oneshot", description="自动探测任务类型(oneshot/periodic)"),
    'auto_group_policy_id': fields.String(required=False, default="", description="自动探测使用的扫描策略ID"),
    'auto_group_interval_hours': fields.Integer(required=False, default=24, description="周期探测运行间隔(小时)"),
    'group_id': fields.String(required=False, default="", description="建组时关联的所属集团ID")
})

def build_auto_group_intent(args):
    """
    从请求参数中提取「自动纳管建组」意图字段。

    该意图由后端调度器 icp_auto_group_scheduler 消费，取代原先依赖浏览器页面轮询的实现。
    """
    try:
        interval = int(args.get('auto_group_interval_hours') or 24)
    except (ValueError, TypeError):
        interval = 24
    return {
        "auto_group": bool(args.get('auto_group', False)),
        "auto_group_scan": bool(args.get('auto_group_scan', False)),
        "auto_group_task_type": args.get('auto_group_task_type') or "oneshot",
        "auto_group_policy_id": args.get('auto_group_policy_id') or "",
        "auto_group_interval_hours": interval,
        "group_id": (args.get('group_id') or "").strip()
    }


@ns.route('/task')
class IcpTask(ARLResource):
    parser = get_arl_parser(search_icp_task_fields, location='args')

    def build_db_query(self, args):
        status_val = args.get("status")
        # 聚合态筛选：与统计口径保持一致（active = running + waiting，abnormal = error + stop），
        # 避免指标卡展示数量与点击筛选后的结果集数量不一致。
        aggregate_status = {
            "active": {"$nin": ["done", "error", "stop"]},
            "abnormal": {"$in": ["error", "stop"]},
            "running": {"$nin": ["done", "error", "stop", "waiting"]},
        }
        if status_val in aggregate_status:
            args_copy = dict(args)
            args_copy.pop("status", None)
            q = super().build_db_query(args_copy)
            q["status"] = aggregate_status[status_val]
            return q
        return super().build_db_query(args)

    @auth
    @ns.expect(parser)
    def get(self):
        """
        获取 ICP 任务列表
        """
        args = self.parser.parse_args()
        include_statistics = args.pop('include_statistics', 0) == 1
        data = self.build_data(args=args, collection='icp_task')
        items = data.get('items', [])
        if items:
            scope_col = conn_db('asset_scope')
            scope_names = [t.get('synced_scope_name') or t.get('name') for t in items if t.get('name')]
            scope_ids = []
            for t in items:
                sid = t.get('synced_scope_id')
                if sid:
                    try:
                        scope_ids.append(bson.ObjectId(sid))
                    except Exception:
                        pass

            query_or = []
            if scope_ids:
                query_or.append({"_id": {"$in": scope_ids}})
            if scope_names:
                query_or.append({"name": {"$in": scope_names}})

            matched_scopes = []
            if query_or:
                matched_scopes = list(scope_col.find(
                    {"$or": query_or},
                    {"_id": 1, "name": 1, "group_name": 1, "domain_array": 1, "scope_array": 1}
                ))

            scope_by_id = {str(s["_id"]): s for s in matched_scopes}
            scope_by_name = {s["name"]: s for s in matched_scopes}

            for item in items:
                stat = item.get('statistic') or {}
                web_cnt = stat.get('web_cnt', 0)

                # 1. 检查是否完全无网站资产
                if web_cnt == 0:
                    item['sync_badge_status'] = 'no_web'
                    item['has_increment'] = False
                    continue

                # 2. 匹配对应资产组（优先按已存储的 synced_scope_id，次选按同名匹配兜底自愈）
                scope = None
                if item.get('synced_scope_id'):
                    scope = scope_by_id.get(str(item['synced_scope_id']))
                if not scope and item.get('name'):
                    scope = scope_by_name.get(item['name'])
                    if scope:
                        # 动态装配返回字段，即刻展示关联状态
                        item['synced_scope_id'] = str(scope['_id'])
                        item['synced_scope_name'] = scope.get('name', '')
                        item['synced_group_name'] = scope.get('group_name', '')

                if scope:
                    item['sync_badge_status'] = 'synced'
                    scope_domains = set((scope.get('domain_array') or scope.get('scope_array') or []))
                    item['has_increment'] = bool(web_cnt > len(scope_domains))
                else:
                    item['sync_badge_status'] = 'unsynced'
                    item['has_increment'] = False

        if include_statistics:
            attach_task_statistics(data, conn_db('icp_task'), logger)

        return data

    @auth
    @ns.expect(add_icp_task_fields)
    def post(self):
        """
        新建企业全景测绘/ICP 查询任务 (支持企业名称自动天眼查全景升级)
        """
        args = self.parse_args(add_icp_task_fields)
        name = args.get('name')
        target = args.get('target')
        query_type = args.get('query_type')

        depth_val = args.get('depth')
        try:
            depth = min(20, max(1, int(depth_val))) if depth_val is not None else 1
        except (ValueError, TypeError):
            depth = 1

        invest_ratio_val = args.get('invest_ratio')
        try:
            invest_ratio = float(invest_ratio_val) if invest_ratio_val is not None else 50.0
        except (ValueError, TypeError):
            invest_ratio = 50.0

        enable_icp = args.get('enable_icp')
        if enable_icp is None:
            enable_icp = True
        else:
            enable_icp = bool(enable_icp)

        target_str = str(target or '').strip()
        custom_name = str(name or '').strip()

        # 尝试通过天眼查 API 进行智能识别与全景升级
        matched_tyc = None
        from app.services.tycClient import TycClient
        try:
            client = TycClient()
            if client.gid and client.token:
                if target_str.startswith("TYC_"):
                    raw_gid = target_str.replace("TYC_", "").strip()
                    cname = client.get_company_name(raw_gid)
                    matched_tyc = {"gid": raw_gid, "company_name": cname or target_str}
                elif target_str.isdigit() and len(target_str) >= 6:
                    cname = client.get_company_name(target_str)
                    matched_tyc = {"gid": target_str, "company_name": cname or target_str}
                else:
                    matched_tyc = client.search_suggest(target_str)
        except Exception as e:
            logger.warning(f"Failed to auto-suggest tyc for target {target_str}: {e}")

        conn = conn_db('icp_task')

        # 分支 1: 成功匹配到天眼查企业 -> 全景双轨任务升级
        if matched_tyc and matched_tyc.get("gid"):
            gid = str(matched_tyc["gid"])
            official_company_name = matched_tyc.get("company_name") or target_str
            final_name = custom_name or (f"{official_company_name}全景测绘" if official_company_name else f"TYC_{gid}")
            full_query_type = ['invest', 'web', 'app', 'mapp', 'wechat', 'weibo']

            task_data = {
                "name": final_name,
                "company_name": official_company_name,
                "target": f"TYC_{gid}",
                "gid": gid,
                "depth": depth,
                "invest_ratio": invest_ratio,
                "query_type": full_query_type,
                "enable_icp": enable_icp,
                "task_type": "tyc",
                "status": TaskStatus.WAITING,
                "start_time": curr_date(),
                "end_time": "-",
                "statistic": {
                    "asset_cnt": 0
                }
            }
            task_data.update(build_auto_group_intent(args))
            result = conn.insert_one(task_data)
            task_id = str(result.inserted_id)

            options = {
                "task_id": task_id,
                "target": f"TYC_{gid}",
                "gid": gid,
                "depth": depth,
                "invest_ratio": invest_ratio,
                "query_type": full_query_type,
                "enable_icp": enable_icp,
                "company_name": official_company_name,
                "name": final_name,
                "type": "tyc",
                "tyc_id": client.gid,
                "tyc_token": client.token
            }
            try:
                _request_recon("start", options)
            except Exception as e:
                logger.error(f"Failed to trigger osint-service for tyc task: {e}")
                conn.update_one({"_id": result.inserted_id, "status": TaskStatus.WAITING}, {"$set": {"status": "error", "error_msg": "请求 osint-service 失败或超时"}})
                return build_ret(ErrorMsg.Error, {"task_id": task_id, "error": "微服务未确认任务派发，请核对任务状态"})

            return build_ret(ErrorMsg.Success, {"task_id": task_id, "matched_tyc": True})

        # 分支 2: 未配置天眼查或未匹配到 -> 工信部全维度官方备案直连
        final_query_type = query_type if (query_type and len(query_type) > 1) else ['web', 'app', 'mapp', 'kapp']
        final_name = custom_name or target_str

        task_data = {
            "name": final_name,
            "target": target_str,
            "query_type": final_query_type,
            "task_type": "icp",
            "status": TaskStatus.WAITING,
            "start_time": curr_date(),
            "end_time": "-",
            "statistic": {
                "asset_cnt": 0
            }
        }
        task_data.update(build_auto_group_intent(args))
        result = conn.insert_one(task_data)
        task_id = str(result.inserted_id)

        options = {
            "task_id": task_id,
            "target": target_str,
            "query_type": final_query_type,
            "type": "icp"
        }
        try:
            _request_recon("start", options)
        except Exception as e:
            logger.error(f"Failed to trigger osint-service: {e}")
            conn.update_one({"_id": result.inserted_id, "status": TaskStatus.WAITING}, {"$set": {"status": "error", "error_msg": "请求 osint-service 失败或超时"}})
            return build_ret(ErrorMsg.Error, {"task_id": task_id, "error": "微服务未确认任务派发，请核对任务状态"})

        return build_ret(ErrorMsg.Success, {"task_id": task_id, "matched_tyc": False})


@ns.route('/task/statistics')
@ns.route('/task/statistics/')
@ns.route('/task/overview_count')
@ns.route('/task/overview_count/')
class IcpTaskStatistics(ARLResource):
    @auth
    def get(self):
        """仅返回企业查询任务的全局状态数量。"""
        try:
            return build_ret(ErrorMsg.Success, get_task_statistics(conn_db('icp_task')))
        except Exception:
            logger.exception('Failed to aggregate enterprise task status statistics')
            return build_ret(ErrorMsg.Error, {'error': '获取企业查询统计失败'})


# 向下兼容别名引用
IcpTaskOverviewCount = IcpTaskStatistics


@ns.route('/tyc_check')
class TycCheck(ARLResource):
    @auth
    def get(self):
        """
        校验天眼查配置是否有效
        """
        from app.services.tycClient import TycClient
        try:
            client = TycClient()
            success, msg = client.check_token()
            return build_ret(ErrorMsg.Success, {"valid": success, "message": msg})
        except Exception as e:
            return build_ret(ErrorMsg.Success, {"valid": False, "message": f"校验异常: {e}"})


@ns.route('/tyc_task')
class TycTask(ARLResource):
    @auth
    @ns.expect(add_tyc_task_fields)
    def post(self):
        """
        新建企业资产查询任务(TYC)
        """
        args = self.parse_args(add_tyc_task_fields)
        name = args.get('name')
        company_name_arg = args.get('company_name')
        gid = args.get('gid')
        depth = min(20, max(1, int(args.get('depth', 1) or 1)))
        invest_ratio = args.get('invest_ratio', 0)
        query_type = args.get('query_type')
        enable_icp = args.get('enable_icp', True)
        if enable_icp is None:
            enable_icp = True

        official_company_name = (company_name_arg or "").strip()

        # 校验天眼查配置是否有效
        from app.services.tycClient import TycClient
        try:
            client = TycClient()
            success, msg = client.check_token()
            if not success:
                return build_ret(msg, {})
        except Exception as e:
            return build_ret(f"校验天眼查配置异常: {e}", {})

        # 若 gid 为企业名称而非纯数字ID，自动通过 search_suggest 提取真实 GID
        gid_str = str(gid or '').strip()
        if not gid_str.isdigit() and len(gid_str) >= 2:
            try:
                suggest_res = client.search_suggest(gid_str)
                if suggest_res and suggest_res.get("gid"):
                    gid = str(suggest_res["gid"])
                    if not official_company_name:
                        official_company_name = suggest_res.get("company_name", "")
            except Exception as e:
                logger.warning(f"Failed to auto-suggest tyc for {gid_str}: {e}")

        # 定向解析企业官方工商全称（彻底脱钩用户任务自定义名称）
        if not official_company_name:
            try:
                official_company_name = client.get_company_name(gid)
            except Exception as e:
                logger.warning(f"Failed to resolve company name for GID {gid}: {e}")

        if not name or not name.strip():
            name = f"{official_company_name}企业测绘" if official_company_name else f"TYC_{gid}"
        else:
            name = name.strip()

        task_data = {
            "name": name,
            "company_name": official_company_name,
            "target": f"TYC_{gid}",
            "gid": gid,
            "depth": depth,
            "invest_ratio": invest_ratio,
            "query_type": query_type,
            "enable_icp": bool(enable_icp),
            "task_type": "tyc",  # 用于前端区分
            "status": TaskStatus.WAITING,
            "start_time": curr_date(),
            "end_time": "-",
            "statistic": {
                "asset_cnt": 0
            },
        }
        # 自动纳管建组意图：由后端调度器消费，不再依赖前端页面轮询
        task_data.update(build_auto_group_intent(args))

        conn = conn_db('icp_task')
        result = conn.insert_one(task_data)
        task_id = str(result.inserted_id)

        options = {
            "task_id": task_id,
            "target": f"TYC_{gid}",
            "gid": gid,
            "depth": depth,
            "invest_ratio": invest_ratio,
            "query_type": query_type,
            "enable_icp": bool(enable_icp),
            "company_name": official_company_name,
            "name": name
        }

        options["type"] = "tyc"
        options["tyc_id"] = client.gid
        options["tyc_token"] = client.token
        try:
            _request_recon("start", options)
        except Exception as e:
            logger.error(f"Failed to trigger osint-service: {e}")
            conn.update_one({"_id": result.inserted_id, "status": TaskStatus.WAITING}, {"$set": {"status": "error", "error_msg": "请求 osint-service 失败或超时"}})
            return build_ret(ErrorMsg.Error, {"task_id": task_id, "error": "微服务未确认任务派发，请核对任务状态"})

        return build_ret(ErrorMsg.Success, {"task_id": task_id})


# 资产搜索字段
base_search_icp_asset_fields = {
    'task_id': fields.String(description="任务ID"),
    'unitName': fields.String(description="主办单位名称"),
    'natureName': fields.String(description="单位性质(ICP)"),
    'domain': fields.String(description="域名"),
    'mainLicence': fields.String(description="主备案号"),
    'serviceLicence': fields.String(description="服务备案号(APP/小程序/快应用)"),
    'contentTypeName': fields.String(description="前置审批/内容类型"),
    'companyName': fields.String(description="TYC主办单位"),
    'ym': fields.String(description="TYC域名"),
    'liscense': fields.String(description="TYC主备案号"),
    'webName': fields.String(description="TYC网站名称"),
    'serviceName': fields.String(description="TYC小程序名称"),
    'serviceFilingNumber': fields.String(description="TYC小程序备案号"),
    'name': fields.String(description="TYC名称(APP/投资/微博)"),
    'classes': fields.String(description="TYC分类"),
    'type': fields.String(description="TYC应用类型"),
    'legalPersonName': fields.String(description="TYC法定代表人"),
    'title': fields.String(description="TYC公众号名称"),
    'publicNum': fields.String(description="TYC公众号"),
    'companyType': fields.String(description="TYC单位性质"),
    'brief': fields.String(description="TYC简介"),
    'status': fields.String(description="TYC企业状态"),
    'recommend': fields.String(description="TYC公众号简介"),
    'href': fields.String(description="TYC微博链接"),
    'amount_num': fields.Float(description="TYC注册资本数字(等于)"),
    'amount_num__ngt': fields.Float(description="TYC注册资本数字(大于)"),
    'amount_num__nlt': fields.Float(description="TYC注册资本数字(小于)"),
    'percent_num': fields.Float(description="TYC投资比例数字(等于)"),
    'percent_num__ngt': fields.Float(description="TYC投资比例数字(大于)"),
    'percent_num__nlt': fields.Float(description="TYC投资比例数字(小于)"),
    'amount': fields.String(description="TYC注册资本(模糊搜索)"),
    'percent': fields.String(description="TYC投资比例(模糊搜索)"),
    'examineDate': fields.String(description="TYC审核日期"),
    'updateRecordTime': fields.String(description="ICP审核日期"),
    'query_type': fields.String(description="查询类型"),
    '_id': fields.String(description="资产ID")
}
base_search_icp_asset_fields.update(base_query_fields)
search_icp_asset_fields = ns.model('SearchIcpAsset', base_search_icp_asset_fields)

def normalize_icp_asset_item(item):
    """
    跨数据源 (ICP vs TYC) 资产字段归一化增强
    保留原始字段完整性，同时补充统一的规范别名字段
    """
    if not isinstance(item, dict):
        return item

    # 1. 域名归一化
    domain = item.get('domain') or item.get('ym') or ''
    item['domain'] = domain
    if 'ym' not in item or not item['ym']:
        item['ym'] = domain

    # 2. 名称归一化 (网站名称/应用名称/公众号标题)
    service_name = item.get('serviceName') or item.get('webName') or item.get('name') or item.get('title') or ''
    item['serviceName'] = service_name
    if not item.get('name'):
        item['name'] = service_name
    if not item.get('webName'):
        item['webName'] = service_name
    if not item.get('title'):
        item['title'] = service_name

    # 3. 主办单位归一化
    unit_name = item.get('unitName') or item.get('companyName') or ''
    if not unit_name and isinstance(item.get('miniProgramIcpRecordDetail'), dict):
        unit_name = item['miniProgramIcpRecordDetail'].get('icpFilingSubjectInformation', {}).get('organizingName', '')
    item['unitName'] = unit_name
    if not item.get('companyName'):
        item['companyName'] = unit_name

    # 4. 备案号归一化
    licence = item.get('serviceLicence') or item.get('liscense') or item.get('serviceFilingNumber') or item.get('mainLicence') or ''
    item['serviceLicence'] = licence
    if not item.get('liscense'):
        item['liscense'] = licence
    if not item.get('serviceFilingNumber'):
        item['serviceFilingNumber'] = licence

    # 5. 更新/审核时间归一化
    update_time = item.get('updateRecordTime') or item.get('examineDate') or ''
    item['updateRecordTime'] = update_time
    if not item.get('examineDate'):
        item['examineDate'] = update_time

    # 6. 首页网址归一化
    home_url = item.get('homeUrl')
    if not home_url:
        ws = item.get('webSite')
        if isinstance(ws, list) and ws:
            home_url = ws[0]
        elif isinstance(ws, str):
            home_url = ws
    item['homeUrl'] = home_url or ''

    # 7. 图标/图片/二维码归一化
    icon = item.get('icon') or item.get('titleImgURL') or item.get('codeImg') or item.get('ico') or item.get('productLogo') or ''
    item['icon'] = icon

    # 8. 简介/描述/认证信息归一化
    brief = item.get('brief') or item.get('recommend') or item.get('info') or ''
    item['brief'] = brief

    # 9. 分类/类型归一化
    category = item.get('category') or item.get('classes') or item.get('type') or ''
    item['category'] = category

    # 10. 微信公众号 ID 归一化
    wechat_id = item.get('wechatId') or item.get('publicNum') or ''
    item['wechatId'] = wechat_id

    # 11. 法人归一化
    legal_person = item.get('legalPerson') or item.get('legalPersonName') or ''
    item['legalPerson'] = legal_person

    # 12. 状态归一化
    status = item.get('status') or item.get('regStatus') or ''
    item['status'] = status

    # 13. 地区归一化
    region = item.get('region') or item.get('province') or ''
    if item.get('city') and item.get('city') not in region:
        region = f"{region} {item.get('city')}".strip()
    item['region'] = region

    return item


@ns.route('/asset')
class IcpAsset(ARLResource):
    parser = get_arl_parser(search_icp_asset_fields, location='args')

    def build_db_query(self, args=None):
        query = super().build_db_query(args)

        # 针对跨数据源 (ICP vs TYC) 的核心字段构建多路兼容 ($or) 查询
        or_groups = []

        # 1. 域名 (domain vs ym)
        if 'domain' in query:
            val = query.pop('domain')
            query.pop('ym', None)
            or_groups.append([{'domain': val}, {'ym': val}])
        elif 'ym' in query:
            val = query.pop('ym')
            or_groups.append([{'domain': val}, {'ym': val}])

        # 2. 单位名称 (unitName vs companyName)
        if 'unitName' in query:
            val = query.pop('unitName')
            query.pop('companyName', None)
            or_groups.append([{'unitName': val}, {'companyName': val}])
        elif 'companyName' in query:
            val = query.pop('companyName')
            or_groups.append([{'unitName': val}, {'companyName': val}])

        # 3. 业务/资产名称 (serviceName, webName, name, title)
        matched_name_val = None
        for k in ['serviceName', 'name', 'webName', 'title']:
            if k in query:
                matched_name_val = query.pop(k)
                break
        if matched_name_val:
            for k in ['serviceName', 'name', 'webName', 'title']:
                query.pop(k, None)
            or_groups.append([
                {'serviceName': matched_name_val},
                {'name': matched_name_val},
                {'webName': matched_name_val},
                {'title': matched_name_val}
            ])

        # 4. 备案号 (serviceLicence, liscense, serviceFilingNumber, mainLicence)
        matched_lic_val = None
        for k in ['serviceLicence', 'liscense', 'serviceFilingNumber', 'mainLicence']:
            if k in query:
                matched_lic_val = query.pop(k)
                break
        if matched_lic_val:
            for k in ['serviceLicence', 'liscense', 'serviceFilingNumber', 'mainLicence']:
                query.pop(k, None)
            or_groups.append([
                {'serviceLicence': matched_lic_val},
                {'liscense': matched_lic_val},
                {'serviceFilingNumber': matched_lic_val},
                {'mainLicence': matched_lic_val}
            ])

        # 5. 法定代表人 (legalPerson vs legalPersonName)
        if 'legalPerson' in query:
            val = query.pop('legalPerson')
            query.pop('legalPersonName', None)
            or_groups.append([{'legalPerson': val}, {'legalPersonName': val}])
        elif 'legalPersonName' in query:
            val = query.pop('legalPersonName')
            or_groups.append([{'legalPerson': val}, {'legalPersonName': val}])

        # 6. 企业状态 (status vs regStatus)
        if 'status' in query:
            val = query.pop('status')
            query.pop('regStatus', None)
            or_groups.append([{'status': val}, {'regStatus': val}])
        elif 'regStatus' in query:
            val = query.pop('regStatus')
            or_groups.append([{'status': val}, {'regStatus': val}])

        # 7. 分类/性质 (category vs classes vs type vs companyType vs natureName)
        matched_cat_val = None
        for k in ['category', 'classes', 'type', 'companyType', 'natureName']:
            if k in query:
                matched_cat_val = query.pop(k)
                break
        if matched_cat_val:
            for k in ['category', 'classes', 'type', 'companyType', 'natureName']:
                query.pop(k, None)
            or_groups.append([
                {'category': matched_cat_val},
                {'classes': matched_cat_val},
                {'type': matched_cat_val},
                {'companyType': matched_cat_val},
                {'natureName': matched_cat_val}
            ])

        # 8. 微信公号 (wechatId vs publicNum)
        if 'wechatId' in query:
            val = query.pop('wechatId')
            query.pop('publicNum', None)
            or_groups.append([{'wechatId': val}, {'publicNum': val}])
        elif 'publicNum' in query:
            val = query.pop('publicNum')
            or_groups.append([{'wechatId': val}, {'publicNum': val}])

        # 9. 简介 (brief vs recommend vs info)
        matched_brief_val = None
        for k in ['brief', 'recommend', 'info']:
            if k in query:
                matched_brief_val = query.pop(k)
                break
        if matched_brief_val:
            for k in ['brief', 'recommend', 'info']:
                query.pop(k, None)
            or_groups.append([
                {'brief': matched_brief_val},
                {'recommend': matched_brief_val},
                {'info': matched_brief_val}
            ])

        if or_groups:
            existing_and = query.get('$and', [])
            for grp in or_groups:
                existing_and.append({'$or': grp})
            query['$and'] = existing_and

        return query

    @auth
    @ns.expect(parser)
    def get(self):
        """
        获取 ICP 资产列表
        """
        args = self.parser.parse_args()
        data = self.build_data(args=args, collection='icp_asset')
        items = data.get('items', [])
        for item in items:
            normalize_icp_asset_item(item)
        return data


@ns.route('/stop/<string:task_id>')
class IcpTaskStop(ARLResource):
    @auth
    def get(self, task_id):
        """停止 ICP 任务 (单任务)"""
        try:
            query = {"_id": bson.ObjectId(task_id)}
            update = {"$set": {"status": TaskStatus.STOP, "end_time": curr_date()}}
            conn_db('icp_task').update_one(query, update)
            
            # 及时通知微服务终止正在执行或等待的协程任务释放锁
            try:
                _request_recon("stop", {"task_id": task_id})
            except Exception as e:
                logger.error(f"Failed to stop osint-service task {task_id}: {e}")
                return build_ret(ErrorMsg.Error, {"task_id": task_id, "error": "微服务未确认停止，请稍后重试停止操作"})
                
            return build_ret(ErrorMsg.Success, {"task_id": task_id})
        except Exception as e:
            return build_ret(ErrorMsg.Error, {"error": str(e)})


@ns.route('/task/log/<string:task_id>')
class IcpTaskLog(ARLResource):
    @auth
    def get(self, task_id):
        """获取 ICP / TYC 任务的运行日志"""
        try:
            query = {"task_id": task_id}
            logs = list(conn_db('syslog').find(query, {"_id": 0}).sort("create_time", 1).limit(1000))
            return {
                "code": 200,
                "message": "success",
                "total": len(logs),
                "items": logs
            }
        except Exception as e:
            return build_ret(ErrorMsg.Error, {"error": str(e)})


@ns.route('/delete/')
class IcpTaskBatchDelete(ARLResource):
    @auth
    def post(self):
        """批量删除 ICP 任务"""
        try:
            from flask import request as flask_request
            args = flask_request.json or {}
            task_ids = args.get('task_ids', [])
            if not isinstance(task_ids, list):
                return build_ret(
                    {"code": 400, "message": "task_ids must be a list"},
                    {"error": "task_ids must be a list"}
                )

            if not task_ids:
                return build_ret(ErrorMsg.Success, {"msg": "成功删除 0 个任务"})

            obj_ids = []
            for tid in task_ids:
                try:
                    obj_ids.append(bson.ObjectId(tid))
                except Exception as err:
                    logger.warning(f"Invalid task_id format {tid}: {err}")

            # 严格前置拦截：只要待删除列表中包含任何 running / waiting 任务，立即整体拒绝本次删除操作
            active_tasks = list(conn_db('icp_task').find(
                {"_id": {"$in": obj_ids}, "status": {"$in": [TaskStatus.RUNNING, TaskStatus.WAITING]}},
                {"_id": 1, "status": 1}
            ))
            if active_tasks:
                return build_ret(
                    {"code": 400, "message": "所选任务中存在正在运行或排队中的任务，请先停止后再删除"},
                    {"active_count": len(active_tasks)}
                )

            success_count = 0
            # 停止记录可能来自失败的停止请求，删除前再次确认后台协程已经退出。
            stopped_tasks = conn_db('icp_task').find({"_id": {"$in": obj_ids}, "status": TaskStatus.STOP}, {"_id": 1})
            for stopped in stopped_tasks:
                _request_recon("stop", {"task_id": str(stopped["_id"])})
            for tid in task_ids:
                try:
                    obj_id = bson.ObjectId(tid)
                    # 原子性状态条件删除，规避并发竞态
                    del_res = conn_db('icp_task').delete_one({
                        "_id": obj_id,
                        "status": {"$nin": [TaskStatus.RUNNING, TaskStatus.WAITING]}
                    })
                    del_count = getattr(del_res, 'deleted_count', 0) if del_res else 0
                    if del_count > 0:
                        conn_db('icp_asset').delete_many({"task_id": str(tid)})
                        # 级联解绑资产组悬空指针，杜绝幽灵绑定
                        conn_db('asset_scope').update_many(
                            {"synced_icp_task_id": str(tid)},
                            {"$unset": {"synced_icp_task_id": ""}}
                        )
                        success_count += 1
                except Exception as item_err:
                    logger.error(f"Failed to cascade delete icp_task {tid}: {item_err}")

            return build_ret(ErrorMsg.Success, {"msg": f"成功删除 {success_count} 个任务", "deleted_count": success_count})
        except Exception as e:
            return build_ret(ErrorMsg.Error, {"error": str(e)})

@ns.route('/delete/<string:task_id>')
class IcpTaskDelete(ARLResource):
    @auth
    def get(self, task_id):
        """删除 ICP 任务 (单任务)"""
        try:
            try:
                obj_id = bson.ObjectId(task_id)
            except Exception:
                return build_ret(
                    {"code": 400, "message": "无效的任务 ID 格式"},
                    {"task_id": task_id}
                )

            task = conn_db('icp_task').find_one({"_id": obj_id})
            if not task:
                return build_ret(ErrorMsg.NotFoundTask, {"task_id": task_id})

            # 严格前置拦截：禁止直接删除运行中或排队中的任务
            if task.get("status") in [TaskStatus.RUNNING, TaskStatus.WAITING]:
                return build_ret(
                    {"code": 400, "message": "无法删除正在运行或排队中的任务，请先手动停止后再行删除"},
                    {"task_id": task_id}
                )

            if task.get("status") == TaskStatus.STOP:
                _request_recon("stop", {"task_id": str(task_id)})

            # 原子性状态条件删除
            del_res = conn_db('icp_task').delete_one({
                "_id": obj_id,
                "status": {"$nin": [TaskStatus.RUNNING, TaskStatus.WAITING]}
            })
            del_count = getattr(del_res, 'deleted_count', 0) if del_res else 0
            if del_count > 0:
                conn_db('icp_asset').delete_many({"task_id": str(task_id)})
                conn_db('asset_scope').update_many(
                    {"synced_icp_task_id": str(task_id)},
                    {"$unset": {"synced_icp_task_id": ""}}
                )

            return build_ret(ErrorMsg.Success, {"task_id": task_id, "deleted_count": del_count})
        except Exception as e:
            return build_ret(ErrorMsg.Error, {"error": str(e)})


def _restart_single_icp_task(task_id: str):
    """
    底层核心逻辑：克隆并重启单个 ICP / TYC 任务
    返回: (ok: bool, reason: str, data: dict)
    """
    try:
        task = conn_db('icp_task').find_one({"_id": bson.ObjectId(task_id)})
        if not task:
            return False, "not_found", {"task_id": task_id, "error": "任务不存在"}

        # 防御逻辑：处于运行中或等待中的任务不允许重复重启
        if task.get("status") in [TaskStatus.WAITING, TaskStatus.RUNNING, "running"]:
            return False, "running", {"task_id": task_id, "error": "任务正在运行或等待中"}

        # 克隆旧任务数据
        task.pop('_id', None)
        task['status'] = TaskStatus.WAITING
        task['start_time'] = curr_date()
        task['end_time'] = "-"
        task['statistic'] = {"asset_cnt": 0}
        task['error_msg'] = ""
        task['prev_task_id'] = str(task_id)

        # 插入新任务
        result = conn_db('icp_task').insert_one(task)
        new_task_id = str(result.inserted_id)

        # 若原任务关联了资产组，记录 updating_icp_task_id，保持原有主指针 synced_icp_task_id 不变
        synced_scope_id = task.get('synced_scope_id')
        if not synced_scope_id:
            scope_doc = conn_db('asset_scope').find_one({"synced_icp_task_id": str(task_id)})
            if scope_doc:
                synced_scope_id = str(scope_doc['_id'])
                conn_db('icp_task').update_one(
                    {"_id": result.inserted_id},
                    {"$set": {"synced_scope_id": synced_scope_id}}
                )

        if synced_scope_id:
            try:
                conn_db('asset_scope').update_one(
                    {"_id": bson.ObjectId(synced_scope_id)},
                    {"$set": {"updating_icp_task_id": new_task_id}}
                )
            except Exception as scope_err:
                logger.warning(f"Failed to update asset_scope updating_icp_task_id: {scope_err}")

        # 格式化 query_type 确保为 list
        query_type = task.get("query_type", ["web"])
        if isinstance(query_type, str):
            query_type = [q.strip() for q in query_type.split(',') if q.strip()]

        # 重新发起
        if task.get("task_type") == "tyc":
            from app.services.tycClient import TycClient
            client = TycClient()
            options = {
                "task_id": new_task_id,
                "target": task.get("target") or f"TYC_{task.get('gid')}",
                "gid": task.get("gid"),
                "depth": task.get("depth", 1),
                "invest_ratio": task.get("invest_ratio", 0),
                "query_type": query_type,
                "enable_icp": task.get("enable_icp", True),
                "company_name": task.get("company_name") or "",
                "name": task.get("name", ""),
                "type": "tyc",
                "tyc_id": client.gid,
                "tyc_token": client.token
            }
        else:
            options = {
                "task_id": new_task_id,
                "target": task.get("target"),
                "query_type": query_type,
                "type": "icp"
            }

        try:
            _request_recon("start", options)
        except Exception as e:
            logger.error(f"Failed to trigger osint-service: {e}")
            failed = conn_db('icp_task').update_one(
                {"_id": result.inserted_id, "status": TaskStatus.WAITING},
                {"$set": {"status": "error", "error_msg": "请求 osint-service 失败或超时"}}
            )
            # 响应丢失时任务可能已经运行，只清理确认置为失败的新任务关联。
            if synced_scope_id and failed.matched_count:
                try:
                    conn_db('asset_scope').update_one(
                        {"_id": bson.ObjectId(synced_scope_id), "updating_icp_task_id": new_task_id},
                        {"$unset": {"updating_icp_task_id": ""}}
                    )
                except Exception:
                    pass
            return False, "trigger_failed", {"task_id": task_id, "error": f"微服务调度失败: {e}"}

        return True, "success", {"old_task_id": task_id, "new_task_id": new_task_id, "updating_task_id": new_task_id}
    except Exception as e:
        logger.error(f"Restart task {task_id} failed: {e}")
        return False, "error", {"task_id": task_id, "error": str(e)}


@ns.route('/restart/<string:task_id>')
class IcpTaskRestart(ARLResource):
    @auth
    def get(self, task_id):
        """重启 ICP 任务 (单任务 GET)"""
        return self._restart(task_id)

    @auth
    def post(self, task_id):
        """重启 ICP 任务 (单任务 POST)"""
        return self._restart(task_id)

    def _restart(self, task_id):
        ok, reason, res = _restart_single_icp_task(task_id)
        if ok:
            return build_ret(ErrorMsg.Success, {
                "task_id": res["new_task_id"],
                "old_task_id": res.get("old_task_id"),
                "updating_task_id": res["new_task_id"],
                "status": "waiting"
            })
        if reason == "not_found":
            return build_ret(ErrorMsg.NotFoundTask, {"task_id": task_id})
        if reason == "running":
            return build_ret(ErrorMsg.TaskIsRunning, {"task_id": task_id})
        return build_ret(ErrorMsg.Error, {"error": res.get("error", "重启失败")})


@ns.route('/restart/')
class IcpTaskBatchRestart(ARLResource):
    @auth
    def post(self):
        """批量重启 ICP / TYC 任务"""
        try:
            from flask import request as flask_request
            args = flask_request.json or {}
            task_ids = args.get('task_ids') or args.get('task_id', [])
            if isinstance(task_ids, str):
                task_ids = [task_ids]
            if not isinstance(task_ids, list):
                return build_ret(ErrorMsg.Error, {"error": "task_ids must be a list"})

            restarted_tasks = []
            skipped_tasks = []
            failed_tasks = []

            for tid in task_ids:
                ok, reason, res = _restart_single_icp_task(str(tid))
                if ok:
                    restarted_tasks.append(res)
                elif reason == "running":
                    skipped_tasks.append(tid)
                else:
                    failed_tasks.append({"task_id": tid, "reason": reason, "error": res.get("error")})

            return build_ret(ErrorMsg.Success, {
                "restarted_count": len(restarted_tasks),
                "skipped_count": len(skipped_tasks),
                "failed_count": len(failed_tasks),
                "restarted_tasks": restarted_tasks,
                "skipped_tasks": skipped_tasks,
                "failed_tasks": failed_tasks
            })
        except Exception as e:
            return build_ret(ErrorMsg.Error, {"error": str(e)})


from flask import make_response

def _build_excel_response(assets, filename="export.xlsx", target_query_type=None):
    try:
        import numpy
        if not hasattr(numpy, "float"):
            numpy.float = float
        if not hasattr(numpy, "int"):
            numpy.int = int
        if not hasattr(numpy, "bool"):
            numpy.bool = bool
    except Exception:
        pass

    import openpyxl
    from io import BytesIO
    from flask import make_response

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # 移除默认的 Sheet

    headers = {
        'web': (['主办单位名称', '单位性质', '主备案号', '域名', '网站名称', '服务许可/备案号', '更新时间/审核日期', '首页网址'],
                ['unitName|companyName', 'natureName|companyType', 'mainLicence|liscense', 'domain|ym', 'serviceName|webName', 'serviceLicence|liscense', 'updateRecordTime|examineDate', 'homeUrl|webSite']),
        'app': (['APP名称', '主办单位名称', '分类/单位性质', '主备案号/应用类型', 'APP备案号', '前置审批/内容类型', '审核时间', '简介'],
                ['name|serviceName', 'unitName|companyName', 'natureName|classes|category', 'mainLicence|type', 'serviceLicence', 'contentTypeName', 'updateRecordTime|examineDate', 'brief|info']),
        'invest': (['投资公司名称', '法定代表人', '注册资本', '投资比例(%)', '企业状态', '地区'],
                ['name', 'legalPersonName|legalPerson', 'amount', 'percent', 'regStatus|status', 'region|province']),
        'wechat': (['公众号名称', '微信号', '简介', '认证主体'],
                ['title|name', 'publicNum|wechatId', 'recommend|brief', 'unitName|companyName']),
        'weibo': (['微博名称', '微博链接', '认证信息/简介', '粉丝数'],
                ['name', 'href', 'info|brief', 'fans']),
        'mapp': (['小程序名称', '主办单位名称', '主备案号', '小程序备案号', '前置审批/内容类型', '审核时间', '简介'],
                ['name|serviceName', 'unitName|companyName', 'mainLicence', 'serviceLicence|serviceFilingNumber', 'contentTypeName', 'updateRecordTime|examineDate', 'brief|recommend']),
        'kapp': (['快应用名称', '主办单位名称', '主备案号', '快应用备案号', '前置审批/内容类型', '审核时间', '简介'],
                ['name|serviceName', 'unitName|companyName', 'mainLicence', 'serviceLicence', 'contentTypeName', 'updateRecordTime|examineDate', 'brief']),
    }

    grouped_assets = {}
    for item in assets:
        normalize_icp_asset_item(item)
        qt = item.get('query_type', 'unknown')
        if target_query_type and target_query_type != 'all' and qt != target_query_type:
            continue
        if qt not in grouped_assets:
            grouped_assets[qt] = []
        grouped_assets[qt].append(item)

    for qt, items in grouped_assets.items():
        if qt in headers:
            h_labels, h_keys = headers[qt]
        else:
            h_labels = ['数据名称']
            h_keys = ['name']

        ws = wb.create_sheet(title=qt)
        append_literal_row(ws, h_labels)
        for item in items:
            row = []
            for k in h_keys:
                if '|' in k:
                    parts = k.split('|')
                    val = ''
                    for p in parts:
                        v = item.get(p)
                        if v is not None and v != '':
                            if isinstance(v, list):
                                val = ', '.join(str(x) for x in v)
                            else:
                                val = v
                            break
                else:
                    val = item.get(k, '')
                row.append(str(val if val is not None else ''))
            append_literal_row(ws, row)

    if len(wb.sheetnames) == 0:
        ws = wb.create_sheet(title='Empty')
        append_literal_row(ws, ['无数据'])

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    response = make_response(output.getvalue())
    response.headers["Content-Disposition"] = f"attachment; filename={filename}"
    response.headers["Content-Type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return response


@ns.route('/export/<string:task_id>')
class IcpTaskExport(ARLResource):
    @auth
    def get(self, task_id):
        """导出 ICP/TYC 资产 (多表单 Excel)"""
        try:
            from flask import request as flask_request
            qt = flask_request.args.get('query_type')
            query = {"task_id": task_id}
            if qt and qt != 'all':
                query["query_type"] = qt
            assets = list(conn_db('icp_asset').find(query))
            return _build_excel_response(assets, filename=f"export_{task_id}_{qt or 'all'}.xlsx", target_query_type=qt)
        except Exception as e:
            return build_ret(ErrorMsg.Error, {"error": str(e)})


batch_export_fields = ns.model('IcpBatchExport', {
    "task_id": fields.List(fields.String(description="任务 ID"), required=True)
})

@ns.route('/batch_export')
class IcpTaskBatchExport(ARLResource):
    @auth
    @ns.expect(batch_export_fields)
    def post(self):
        """批量导出 ICP/TYC 资产 (多表单 Excel)"""
        try:
            args = self.parse_args(batch_export_fields)
            task_id_list = args.get("task_id", [])
            if not task_id_list:
                return build_ret(ErrorMsg.Error, {"error": "未提供任务 ID 列表"})

            assets = list(conn_db('icp_asset').find({"task_id": {"$in": task_id_list}}))
            return _build_excel_response(assets, filename="batch_export.xlsx")
        except Exception as e:
            return build_ret(ErrorMsg.Error, {"error": str(e)})



def perform_icp_sync(task_id, args):
    """
    企业测绘任务 → 资产组同步核心逻辑（与 Flask 请求上下文解耦）

    供 /icp/sync 接口与后端自动建组调度器共同复用，避免两套业务逻辑漂移。
    :param task_id: icp_task 的 _id 字符串
    :param args: 同步参数字典，字段语义同 /icp/sync 请求体
    :return: (ok: bool, err_code, payload: dict)
    """
    from app.helpers.scope import trigger_scope_domain_scan
    mode = args.get('mode', 'new')  # 'new' or 'existing'
    target_name = (args.get('target_name') or '').strip()
    scope_id = args.get('scope_id')
    group_id = (args.get('group_id') or '').strip()
    selected_domains = args.get('selected_domains')
    auto_scan = args.get('auto_scan', False)
    task_type = args.get('task_type', 'oneshot')
    policy_id = args.get('policy_id', '')
    interval_hours = int(args.get('interval_hours', 24) or 24)

    try:
        task = conn_db('icp_task').find_one({"_id": bson.ObjectId(task_id)})
    except Exception:
        return False, ErrorMsg.Error, {"error": "任务 ID 格式非法"}
    if not task:
        return False, ErrorMsg.NotFoundTask, {"task_id": task_id}

    task_name = task.get('name', 'ICP查询任务')

    # 校验所属集团
    group_name = ""
    if group_id:
        try:
            group = conn_db('asset_group').find_one({"_id": bson.ObjectId(group_id)})
            if group:
                group_name = group.get('name', '')
            else:
                group_id = ""
        except Exception:
            group_id = ""

    # 查出当前任务下所有 web 查询的资产 (只取网站)
    assets = list(conn_db('icp_asset').find({"task_id": task_id, "query_type": "web"}))
    domains = set()
    for asset in assets:
        d = asset.get('domain') or asset.get('ym')
        if d and isinstance(d, str):
            domains.add(d.strip())

    # 若前端指定了勾选的特定域名列表，则精准过滤
    if isinstance(selected_domains, list) and len(selected_domains) > 0:
        selected_set = {str(d).strip().lower() for d in selected_domains if str(d).strip()}
        domains = {d for d in domains if d.lower() in selected_set}

    if not domains:
        return False, ErrorMsg.Error, {"error": "未发现匹配的网站资产，无法同步"}

    final_scope_id = None
    task_triggered_count = 0

    # 封装生成 domain_status 的元数据辅助函数
    def build_domain_meta(existing_meta=None):
        meta = existing_meta if isinstance(existing_meta, dict) else {}
        meta.update({
            "status": meta.get("status", "unprobed"),
            "sync_source": "icp",
            "task_id": str(task["_id"]),
            "task_name": task_name,
            "sync_time": curr_date()
        })
        return meta

    if mode == 'existing' and scope_id:
        # 关联已有资产
        scope = conn_db('asset_scope').find_one({"_id": bson.ObjectId(scope_id)})
        if not scope:
            return False, ErrorMsg.Error, {"error": "指定的资产组不存在"}
        final_scope_id = str(scope["_id"])
        old_array = scope.get("scope_array", [])
            
        # Calculate duplicates and new additions
        old_set = set(old_array)
        duplicates = domains.intersection(old_set)
        duplicate_count = len(duplicates)
        new_additions = domains - old_set
        insert_count = len(new_additions)
            
        new_array = list(old_set | domains)
        old_domain_array = scope.get("domain_array")
        if old_domain_array is None:
            old_domain_array = old_array if scope.get("scope_type", "domain") == "domain" else []
        new_domain_array = list(set(old_domain_array) | domains)

        domain_status = scope.get("domain_status", {})
        if not isinstance(domain_status, dict):
            domain_status = {}

        for d in domains:
            domain_status[d] = build_domain_meta(domain_status.get(d))

        update_fields = {
            "scope_array": new_array,
            "scope": ",".join(new_array),
            "domain_array": new_domain_array,
            "domain_status": domain_status
        }
        # 若已有资产组未分配集团，且同步时指定了集团，则补充绑定
        if group_id and not scope.get('group_id'):
            update_fields["group_id"] = group_id
            update_fields["group_name"] = group_name

        conn_db('asset_scope').update_one(
            {"_id": scope["_id"]},
            {"$set": update_fields}
        )
        target_name = scope.get('name')
        if not group_name:
            group_name = scope.get('group_name', '')
    else:
        # 新建资产组
        if not target_name:
            target_name = task_name

        # 检查是否同名，不再限制 scope_type
        scope = conn_db('asset_scope').find_one({"name": target_name})
        if scope:
            final_scope_id = str(scope["_id"])
            old_array = scope.get("scope_array", [])
                
            old_set = set(old_array)
            duplicates = domains.intersection(old_set)
            duplicate_count = len(duplicates)
            new_additions = domains - old_set
            insert_count = len(new_additions)

            new_array = list(old_set | domains)
            old_domain_array = scope.get("domain_array")
            if old_domain_array is None:
                old_domain_array = old_array if scope.get("scope_type", "domain") == "domain" else []
            new_domain_array = list(set(old_domain_array) | domains)

            domain_status = scope.get("domain_status", {})
            if not isinstance(domain_status, dict):
                domain_status = {}

            for d in domains:
                domain_status[d] = build_domain_meta(domain_status.get(d))

            update_fields = {
                "scope_array": new_array,
                "scope": ",".join(new_array),
                "domain_array": new_domain_array,
                "domain_status": domain_status
            }
            if group_id and not scope.get('group_id'):
                update_fields["group_id"] = group_id
                update_fields["group_name"] = group_name

            conn_db('asset_scope').update_one(
                {"_id": scope["_id"]},
                {"$set": update_fields}
            )
            if not group_name:
                group_name = scope.get('group_name', '')
        else:
            new_array = list(domains)
            new_additions = set(domains)
            insert_count = len(domains)
            duplicate_count = 0
                
            domain_status = {}
            for d in domains:
                domain_status[d] = build_domain_meta()

            scope_data = {
                "name": target_name,
                "group_id": group_id,
                "group_name": group_name,
                "scope_type": "mixed",
                "scope": ",".join(new_array),
                "scope_array": new_array,
                "domain_array": new_array,
                "domain_status": domain_status,
                "ip_array": []
            }
            insert_res = conn_db('asset_scope').insert_one(scope_data)
            final_scope_id = str(insert_res.inserted_id)

    # 回写 icp_task 最近同步状态
    conn_db('icp_task').update_one(
        {"_id": task["_id"]},
        {"$set": {
            "synced_scope_id": final_scope_id,
            "synced_scope_name": target_name,
            "synced_group_name": group_name,
            "synced_at": curr_date()
        }}
    )

    # 如果用户勾选了自动下发任务且存在新增域名
    if auto_scan and policy_id and new_additions and final_scope_id:
        try:
            task_triggered_count = trigger_scope_domain_scan(
                scope_id=final_scope_id,
                domains=list(new_additions),
                policy_id=policy_id,
                task_type=task_type,
                interval=interval_hours * 3600
            )
        except Exception as e:
            logger.error(f"Auto trigger scan error after sync: {e}")

    return True, ErrorMsg.Success, {
        "msg": f"成功同步了 {len(domains)} 个域名到资产分组 '{target_name}'",
        "insert_count": insert_count,
        "duplicate_count": duplicate_count,
        "task_triggered_count": task_triggered_count,
        "new_domains": list(new_additions),
        "target_name": target_name,
        "scope_id": final_scope_id,
        "group_id": group_id,
        "group_name": group_name,
        "mode": mode
    }


def build_auto_group_payload(task, domains):
    """
    由已完成的企业测绘任务推导自动建组参数。
    与前端向导保持同一语义：新建资产组、名称沿用企业官方名称、按需联动主动探测。
    """
    return {
        "mode": "new",
        "target_name": _resolve_auto_group_name(task),
        "group_id": task.get("group_id") or "",
        "selected_domains": domains,
        "auto_scan": bool(task.get("auto_group_scan")),
        "task_type": task.get("auto_group_task_type") or "oneshot",
        "policy_id": task.get("auto_group_policy_id") or "",
        "interval_hours": int(task.get("auto_group_interval_hours") or 24)
    }


def _resolve_auto_group_name(task):
    """推导自动建组使用的资产组名称，优先企业官方名称"""
    company_name = (task.get("company_name") or "").strip()
    if company_name:
        return company_name
    name = (task.get("name") or "").strip()
    if name:
        # 去掉测绘任务名后缀，还原企业名。
        # 必须按长度降序匹配，否则「企业全景测绘」会先命中较短的「全景测绘」而残留「企业」。
        for suffix in sorted(("企业全景测绘", "企业测绘", "全景测绘"), key=len, reverse=True):
            if name.endswith(suffix):
                name = name[: -len(suffix)].strip()
                break
        if name:
            return name
    return task.get("target") or "企业资产组"


@ns.route('/sync/<string:task_id>')
class IcpTaskSync(ARLResource):
    @auth
    def post(self, task_id):
        """将 ICP/TYC 任务中的网站域名同步至资产分组"""
        from flask import request as flask_request
        try:
            args = flask_request.json or {}
            ok, err_code, payload = perform_icp_sync(task_id, args)
            if not ok:
                return build_ret(err_code, payload)
            return build_ret(ErrorMsg.Success, payload)
        except Exception as e:
            logger.exception(e)
            return build_ret(ErrorMsg.Error, {"error": str(e)})

